import os
from typing import List
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .models import Violation, PaperMetadata

def set_cell_background(cell, hex_color: str):
    """Set table cell background shading."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell internal margins in twips (1/20 of a pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def generate_word_report(
    violations: List[Violation], 
    metadata: PaperMetadata, 
    output_path: str,
    conference_name: str = "ICoDMI 2026 (SSMI IPB)"
) -> str:
    """
    Generates a Microsoft Word (.docx) document with the exact revision table requested by the user:
    Columns: No | Committee Comments | Location in Manuscript
    """
    doc = docx.Document()
    
    # Set page margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run(f"{conference_name}")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(24, 43, 73)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(12)
    run_sub = sub_p.add_run("IEEE Manuscript Format Revision Sheet")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    
    # Metadata Info
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    meta_p.add_run("Manuscript Name: ").bold = True
    meta_p.add_run(f"{os.path.basename(metadata.filename)}\n")
    meta_p.add_run("Page Count / Size: ").bold = True
    meta_p.add_run(f"{metadata.total_pages} pages ({metadata.paper_size_name})\n")
    meta_p.add_run("Total Items to Revise: ").bold = True
    run_status = meta_p.add_run(f"{len(violations)} item(s)")
    if len(violations) > 0:
        run_status.font.color.rgb = RGBColor(180, 40, 40)
        run_status.bold = True
        
    # Revision Table (Matching user's screenshot)
    # Columns: No | Committee Comments | Location in Manuscript
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Table Header Row
    hdr_cells = table.rows[0].cells
    hdr_titles = ["No", "Committee Comments", "Location in Manuscript"]
    col_widths = [Inches(0.6), Inches(4.3), Inches(1.8)]
    
    for i, title in enumerate(hdr_titles):
        cell = hdr_cells[i]
        cell.width = col_widths[i]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(cell, "F2F4F7")
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 30, 30)
        
    # Populate rows
    if not violations:
        row_cells = table.add_row().cells
        for i in range(3):
            row_cells[i].width = col_widths[i]
            set_cell_margins(row_cells[i], top=120, bottom=120, left=140, right=140)
        row_cells[0].paragraphs[0].add_run("-")
        row_cells[1].paragraphs[0].add_run("All checked IEEE format rules are compliant! No revisions needed.")
        row_cells[2].paragraphs[0].add_run("-")
    else:
        for idx, v in enumerate(violations, start=1):
            row = table.add_row()
            cells = row.cells
            for i in range(3):
                cells[i].width = col_widths[i]
                set_cell_margins(cells[i], top=120, bottom=120, left=140, right=140)
                
            # Column 1: No
            p0 = cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(str(idx))
            r0.font.name = "Times New Roman"
            r0.font.size = Pt(10)
            
            # Column 2: Committee Comments
            p1 = cells[1].paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r1_cat = p1.add_run(f"[{v.category}] ")
            r1_cat.font.name = "Times New Roman"
            r1_cat.font.size = Pt(10)
            r1_cat.font.bold = True
            
            r1_comm = p1.add_run(v.comment)
            r1_comm.font.name = "Times New Roman"
            r1_comm.font.size = Pt(10)
            
            # Column 3: Location in Manuscript
            p2 = cells[2].paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r2 = p2.add_run(v.location)
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(10)
            
    # Apply standard table borders
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)
    
    doc.save(output_path)
    return output_path

def generate_excel_report(
    violations: List[Violation], 
    metadata: PaperMetadata, 
    output_path: str
) -> str:
    """Generates an Excel spreadsheet with the same table columns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Review Comments"
    
    # Header styling
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Times New Roman", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    headers = ["No", "Committee Comments", "Location in Manuscript"]
    ws.append(headers)
    
    for col_num in range(1, 4):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_num == 1 else "left", vertical="center")
        
    for idx, v in enumerate(violations, start=1):
        row_data = [idx, f"[{v.category}] {v.comment}", v.location]
        ws.append(row_data)
        current_row = ws.max_row
        for col_num in range(1, 4):
            c = ws.cell(row=current_row, column=col_num)
            c.font = data_font
            c.border = thin_border
            c.alignment = Alignment(
                horizontal="center" if col_num == 1 else "left", 
                vertical="top",
                wrap_text=True
            )
            
    # Set explicit column widths
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 65
    ws.column_dimensions['C'].width = 30
    
    wb.save(output_path)
    return output_path

